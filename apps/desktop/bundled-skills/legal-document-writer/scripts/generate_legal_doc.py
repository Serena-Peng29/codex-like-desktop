#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成符合中国法律文书排版规范的 Word 文档（.docx）。

支持的五类文档（doc_type）：
  complaint   民事起诉状 / 行政起诉状 / 刑事自诉状
  defense     答辩状（民事/行政答辩状）
  judgment    民事判决书 / 判决书
  mediation   民事调解书
  assessment  诉讼评估报告 / 诉讼风险告知

用法：
  python generate_legal_doc.py --spec spec.json --output out.docx

spec.json 结构（顶层字段）：
  doc_type   必填，上述五类之一
  title      必填，文书标题（如 "民事起诉状"）
  subtitle   可选，标题下一行的居中副题（如案号 "（2026）鲁0102民初1234号"）
  font_size  可选，正文字号（磅），默认 16（三号）；诉讼文书常用 12（小四）或 14（四号）
  blocks     必填，正文内容块的有序列表，见下方 block 说明

block 类型（type 字段）：
  heading  标题行  -> level: 1|2|3 控制字体（1=黑体, 2=楷体, 3=仿宋加粗），text 已含编号如 "一、诉讼请求"
  para     正文段  -> text；可选 indent（默认 True 首行缩进2字符）、bold、align（left/center/right/justify）
  center   居中段  -> text；可选 bold
  items    编号列表 -> items: ["1. ...", "2. ..."]，逐条输出为独立段落
  table    表格    -> headers: [...], rows: [[...], ...]，用于对比/清单类内容
  signature 落款   -> lines: ["具状人：张三", "2026年9月2日"]，右对齐
  blank    空行    -> 无额外字段

示例 spec 见 assets/*.json。
"""

import argparse
import json
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ----------------------------- 排版常量 -----------------------------

FONT_TITLE_EA = "宋体"        # 大标题（如需可用方正小标宋）
FONT_HEI = "黑体"             # 一级标题
FONT_KAI = "楷体"             # 二级标题
FONT_FANG = "仿宋_GB2312"     # 正文
FONT_ASCII = "Times New Roman"  # 数字/字母字体

SIZE_TITLE = 22               # 二号
SIZE_BODY_DEFAULT = 16        # 三号
LINE_SPACING_EXACT = 28       # 固定行距（磅）

ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _set_run_font(run, ea=FONT_FANG, size=None, bold=False, ascii_font=FONT_ASCII):
    """设置 run 的中文字体(eastAsia)、西文字体、字号、加粗。"""
    if size is None:
        size = SIZE_BODY_DEFAULT
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), ea)
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)


def _new_paragraph(doc, align=None, line_spacing=LINE_SPACING_EXACT,
                   space_after=0, space_before=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.line_spacing = Pt(line_spacing)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    return p


def add_title(doc, text, subtitle=None):
    p = _new_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    r = p.add_run(text)
    _set_run_font(r, ea=FONT_TITLE_EA, size=SIZE_TITLE, bold=True)
    if subtitle:
        p2 = _new_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
        r2 = p2.add_run(subtitle)
        _set_run_font(r2, ea=FONT_FANG, size=SIZE_BODY_DEFAULT, bold=False)


def add_heading(doc, text, level, body_size):
    ea = {1: FONT_HEI, 2: FONT_KAI, 3: FONT_FANG}.get(level, FONT_HEI)
    bold = level in (1, 3)
    p = _new_paragraph(doc, space_before=6, space_after=0)
    r = p.add_run(text)
    _set_run_font(r, ea=ea, size=body_size, bold=bold)
    return p


def add_para(doc, text, body_size, indent=True, bold=False, align="justify"):
    p = _new_paragraph(doc, align=ALIGN_MAP.get(align, WD_ALIGN_PARAGRAPH.JUSTIFY))
    if indent:
        p.paragraph_format.first_line_indent = Pt(body_size * 2)
    r = p.add_run(text)
    _set_run_font(r, ea=FONT_FANG, size=body_size, bold=bold)
    return p


def add_items(doc, items, body_size):
    for it in items:
        p = _new_paragraph(doc, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.first_line_indent = Pt(body_size * 2)
        r = p.add_run(str(it))
        _set_run_font(r, ea=FONT_FANG, size=body_size)
    return p


def add_table(doc, headers, rows, body_size):
    n_cols = len(headers)
    table = doc.add_table(rows=1, cols=n_cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # 表头
    hdr = table.rows[0].cells
    for j, h in enumerate(headers):
        hdr[j].text = ""
        p = hdr[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(h))
        _set_run_font(r, ea=FONT_HEI, size=body_size, bold=True)
    # 数据行
    for row in rows:
        cells = table.add_row().cells
        for j in range(n_cols):
            val = row[j] if j < len(row) else ""
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(val))
            _set_run_font(r, ea=FONT_FANG, size=body_size)
    return table


def add_signature(doc, lines, body_size):
    for line in lines:
        p = _new_paragraph(doc, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=3)
        r = p.add_run(str(line))
        _set_run_font(r, ea=FONT_FANG, size=body_size)


def build_document(spec):
    doc = Document()

    # 页面设置：A4 + 公文边距（GB/T 9704-2012）
    sec = doc.sections[0]
    sec.page_height = Cm(29.7)
    sec.page_width = Cm(21.0)
    sec.top_margin = Cm(3.7)
    sec.bottom_margin = Cm(3.5)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)

    body_size = int(spec.get("font_size", SIZE_BODY_DEFAULT))

    # 默认样式兜底
    style = doc.styles["Normal"]
    style.font.name = FONT_ASCII
    style.font.size = Pt(body_size)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_FANG)

    add_title(doc, spec.get("title", ""), spec.get("subtitle"))

    for block in spec.get("blocks", []):
        btype = block.get("type", "para")
        if btype == "heading":
            add_heading(doc, block.get("text", ""), int(block.get("level", 1)), body_size)
        elif btype == "para":
            add_para(
                doc, block.get("text", ""), body_size,
                indent=block.get("indent", True),
                bold=block.get("bold", False),
                align=block.get("align", "justify"),
            )
        elif btype == "center":
            p = _new_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
            r = p.add_run(block.get("text", ""))
            _set_run_font(r, ea=FONT_FANG, size=body_size, bold=block.get("bold", False))
        elif btype == "items":
            add_items(doc, block.get("items", []), body_size)
        elif btype == "table":
            add_table(doc, block.get("headers", []), block.get("rows", []), body_size)
        elif btype == "signature":
            add_signature(doc, block.get("lines", []), body_size)
        elif btype == "blank":
            _new_paragraph(doc)
        else:
            sys.stderr.write("未知 block 类型：%s\n" % btype)

    return doc


def main():
    ap = argparse.ArgumentParser(description="生成规范格式的中国法律文书 docx")
    ap.add_argument("--spec", required=True, help="JSON 规格文件路径")
    ap.add_argument("--output", required=True, help="输出 docx 路径")
    args = ap.parse_args()

    with open(args.spec, "r", encoding="utf-8") as f:
        spec = json.load(f)

    out_dir = os.path.dirname(os.path.abspath(args.output))
    if out_dir and not os.path.isdir(out_dir):
        os.makedirs(out_dir, exist_ok=True)

    doc = build_document(spec)
    doc.save(args.output)
    print("已生成：%s" % os.path.abspath(args.output))


if __name__ == "__main__":
    main()
